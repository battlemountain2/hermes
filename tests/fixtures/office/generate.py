import argparse
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.shared import Inches as WordInches, Pt
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.page import PageMargins
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt as SlidePt


def print_layout(sheet, area, landscape=False, height=1):
    sheet.print_area = area
    sheet.sheet_state = "visible"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = height
    sheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.25, bottom=0.25, header=0, footer=0
    )


def generate_workbooks(output):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append(["Item", "Count", "Status"])
    sheet.append(["Generated images", 3, "Ready"])
    sheet.append(["Selectable cells", 12, "Checked"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="244669")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in ["A", "B", "C"]:
        sheet.column_dimensions[column].width = 24
    for column in "DEFGHIJKL":
        sheet.column_dimensions[column].width = 12
    for row in range(1, 23):
        sheet.row_dimensions[row].height = 18
    sheet.merge_cells("A6:C7")
    sheet["A6"] = "Field station totals"
    sheet["A6"].fill = PatternFill("solid", fgColor="244669")
    sheet["A6"].font = Font(name="DejaVu Sans", size=18, bold=True, color="FFFFFF")
    sheet["A6"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("A9:C10")
    sheet["A9"] = "Merged summary"
    sheet["A9"].fill = PatternFill("solid", fgColor="277D5B")
    sheet["A9"].font = Font(name="DejaVu Sans", size=16, bold=True, color="FFFFFF")
    sheet["A9"].alignment = Alignment(horizontal="center", vertical="center")
    for row_number, row in enumerate([
        ("Station", "Visits", "Cadence"),
        ("Aspen", 18, "Weekly"),
        ("Juniper", 27, "Weekly"),
        ("Willow", 12, "Monthly"),
        ("Cedar", 24, "Weekly"),
    ], 12):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_number, column, value)
            cell.fill = PatternFill("solid", fgColor="F4DD79" if row_number % 2 == 0 else "FBEBD9")
            cell.font = Font(name="DejaVu Sans", size=12, color="172D45", bold=row_number == 12)
            cell.border = Border(bottom=Side(style="thin", color="D5DDE5"))
            cell.alignment = Alignment(horizontal="center" if column == 2 else "left")
    chart = BarChart()
    chart.title = "Survey coverage"
    chart.y_axis.title = "Visits"
    chart.x_axis.title = "Station"
    chart.add_data(Reference(sheet, min_col=2, min_row=12, max_row=16), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=13, max_row=16))
    chart.series[0].graphicalProperties.solidFill = "E37434"
    chart.series[0].graphicalProperties.line.solidFill = "E37434"
    chart.legend = None
    chart.width = 18
    chart.height = 7
    sheet.add_chart(chart, "E6")
    print_layout(sheet, "A1:L22", landscape=True)

    details = workbook.create_sheet("Details")
    details.append(["Another sheet", "Keeps navigation"])
    details.merge_cells("A3:D4")
    details["A3"] = "Details page only"
    details["A3"].fill = PatternFill("solid", fgColor="277D5B")
    details["A3"].font = Font(name="DejaVu Sans", size=18, bold=True, color="FFFFFF")
    details["A3"].alignment = Alignment(horizontal="center", vertical="center")
    details.append(["Site notes", "Synthetic observations"])
    details.append(["Marmalade meadow", "Second-sheet navigation verified"])
    details.append(["Read-only source", "No workbook edits during preview"])
    for column in "ABCD":
        details.column_dimensions[column].width = 24
    print_layout(details, "A1:D14")
    workbook.save(output / "layout.xlsx")

    full_workbook = Workbook()
    full_sheet = full_workbook.active
    full_sheet.title = "Full workbook"
    full_sheet.append(["Record", "Count", "Status"])
    for index in range(1, 226):
        label = "Full workbook first page" if index == 1 else f"Synthetic record {index:03d}"
        if index == 225:
            label = "Beyond row 200 sentinel"
        full_sheet.append([label, index, "Generated"])
    for column in "ABC":
        full_sheet.column_dimensions[column].width = 30
    for row in range(1, 227):
        full_sheet.row_dimensions[row].height = 14
    for row in range(46, 226, 45):
        full_sheet.row_breaks.append(Break(id=row))
    full_sheet.print_title_rows = "1:1"
    print_layout(full_sheet, "A1:C226", height=0)
    full_workbook.save(output / "full-workbook.xlsx")


def remove_saved_thumbnail(target):
    relationships_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(target) as archive:
        members = [(item, archive.read(item)) for item in archive.infolist()
                   if not item.filename.startswith("docProps/thumbnail.")]
    with zipfile.ZipFile(target, "w") as archive:
        for item, content in members:
            if item.filename == "_rels/.rels":
                relationships = ElementTree.fromstring(content)
                for relationship in list(relationships):
                    if relationship.get("Type", "").endswith("/metadata/thumbnail"):
                        relationships.remove(relationship)
                ElementTree.register_namespace("", relationships_namespace)
                content = ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True)
            archive.writestr(item, content)


def generate_saved_poster(output):
    poster = Image.new("RGB", (960, 540), "#265185")
    drawing = ImageDraw.Draw(poster)
    font = ImageFont.truetype("DejaVuSans.ttf", 38)
    drawing.text((56, 36), "Hermes slide 1", font=font, fill="#ffffff")
    artwork = Image.open(output / "generated-art.png")
    poster.paste(artwork.resize((400, 200)), (60, 160))
    drawing.rounded_rectangle((520, 160, 880, 360), radius=22, fill="#f4dd79")
    drawing.text((540, 220), "Selectable slide", font=ImageFont.truetype("DejaVuSans.ttf", 24), fill="#172d45")
    drawing.text((540, 254), "text 1", font=font, fill="#172d45")
    poster.save(output / "saved-poster.png")
    relationships_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(output / "layout.pptx") as original:
        with zipfile.ZipFile(output / "saved-poster.pptx", "w") as archive:
            for item in original.infolist():
                content = original.read(item)
                if item.filename == "_rels/.rels":
                    relationships = ElementTree.fromstring(content)
                    ElementTree.SubElement(relationships, f"{{{relationships_namespace}}}Relationship", {
                        "Id": "hermesSavedPoster",
                        "Type": f"{relationships_namespace}/metadata/thumbnail",
                        "Target": "docProps/thumbnail.png",
                    })
                    ElementTree.register_namespace("", relationships_namespace)
                    content = ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True)
                archive.writestr(item, content)
            archive.write(output / "saved-poster.png", "docProps/thumbnail.png")


def generate(output):
    output.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (360, 180), "#e37434")
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((180, 0, 359, 179), fill="#2466b0")
    drawing.ellipse((110, 30, 250, 170), fill="#f4dd79")
    artwork = output / "generated-art.png"
    image.save(artwork)

    document = Document()
    document.styles["Normal"].font.name = "DejaVu Sans"
    document.styles["Normal"].font.size = Pt(12)
    document.add_heading("Hermes layout fixture", 0)
    document.add_paragraph("Selectable first-page text with bold ").add_run("formatting").bold = True
    document.add_picture(str(artwork), width=WordInches(4))
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("Item", "Count")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("Generated shapes", "3")):
        cell.text = value
    document.add_page_break()
    document.add_heading("Second document page", 1)
    document.add_paragraph("Only the second page contains the word marmalade.")
    document.save(output / "layout.docx")

    content = '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" xmlns:xlink="http://www.w3.org/1999/xlink" office:version="1.2">
<office:automatic-styles><style:style style:name="NewPage" style:family="paragraph"><style:paragraph-properties fo:break-before="page"/></style:style></office:automatic-styles>
<office:body><office:text><text:h text:outline-level="1">Hermes ODT fixture</text:h><text:p>Selectable OpenDocument text.</text:p><text:p><draw:frame draw:name="Artwork" text:anchor-type="as-char" svg:width="4in" svg:height="2in"><draw:image xlink:href="Pictures/generated.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></text:p><text:p text:style-name="NewPage">ODT marmalade on the second page.</text:p></office:text></office:body></office:document-content>'''
    manifest = '''<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/><manifest:file-entry manifest:full-path="Pictures/generated.png" manifest:media-type="image/png"/></manifest:manifest>'''
    with zipfile.ZipFile(output / "layout.odt", "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
        archive.writestr("META-INF/manifest.xml", manifest)
        archive.write(artwork, "Pictures/generated.png")

    presentation = Presentation()
    presentation.slide_width = Inches(12)
    presentation.slide_height = Inches(6.75)
    for index, color in enumerate(((38, 81, 133), (39, 125, 91)), 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(*color)
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(10.5), Inches(1))
        paragraph = title.text_frame.paragraphs[0]
        paragraph.text = f"Hermes slide {index}"
        paragraph.font.name = "DejaVu Sans"
        paragraph.font.size = SlidePt(32)
        paragraph.font.color.rgb = RGBColor(255, 255, 255)
        slide.shapes.add_picture(str(artwork), Inches(0.75), Inches(2), width=Inches(5))
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(2), Inches(4.5), Inches(2.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(244, 221, 121)
        shape.text = f"Selectable slide text {index}"
        shape.text_frame.paragraphs[0].font.name = "DejaVu Sans"
        shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(23, 45, 69)
    presentation.save(output / "layout.pptx")

    generate_workbooks(output)

    for filename in ["layout.docx", "layout.pptx"]:
        remove_saved_thumbnail(output / filename)
    generate_saved_poster(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    generate(parser.parse_args().output)
